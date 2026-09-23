#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
noise_scan.py v4 —— 雜訊存活分析（獨立腳本，不碰 run_bot.py）

【v4 和前三版的根本差別】
前三版在算「哪個停損距離的平均毛利最高」。在接近隨機的價格上，那個答案
永遠是「越緊越好」—— 緊的停損虧得少。那是數學必然，不是交易發現，
而且和實盤事實相反（緊貼 0.1% 兩秒就被掃掉）。

v4 改問一個問題，而且只問這一個：
    「要離多遠，部位才不會被日常震盪掃掉？」

做法：從每個時點進場，量價格【第一次】逆向走到 X% 需要多久。
     對每個距離 X，算出撐過 10/30/60/120/300 秒的比例。
     雜訊% = 讓部位能撐過 60 秒、存活率達 70% 的那個距離。

這個定義可以直接驗證：把 WIF 的 0.1% 那一列拉出來看 10 秒存活率，
如果模型是對的，它應該很低 —— 因為你實盤就是兩秒被掃。

用法
  python3 noise_scan.py probe
  python3 noise_scan.py run                          預設 7 天，重新下載
  python3 noise_scan.py run --days 7 --fresh         強制重抓不讀快取
  python3 noise_scan.py run --cache                  讀既有快取（快，但資料是舊的）
  python3 noise_scan.py run --no-mail --only WIFUSDT
"""

import os, sys, time, json, smtplib, datetime
import urllib.request
from email.message import EmailMessage

import numpy as np
from numpy.lib.stride_tricks import sliding_window_view
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

BASE      = "/srv/1111bot"
CACHE_DIR = os.path.join(BASE, "data", "noise_cache")
OUT_DIR   = os.path.join(BASE, "data")
OKX       = "https://www.okx.com"

SYMBOLS = ["AAVEUSD", "ADAUSDT", "AVAXUSDT", "BTCUSDT", "DOGEUSDT", "ETHUSDT",
           "HYPEUSDT", "LINKUSDT", "SOLUSDT", "SUIUSDT", "TAOUSDT", "WIFUSDT",
           "WLDUSDT", "XAUUSDT", "XRPUSDT", "ZECUSDT"]

# 候選距離（%）—— 涵蓋你現在用的 0.1% 到遠得多的 3%
DIST = [0.02, 0.03, 0.05, 0.075, 0.10, 0.15, 0.20, 0.25, 0.30, 0.40,
        0.50, 0.75, 1.00, 1.50, 2.00, 3.00]
SURV_SEC   = [10, 30, 60, 120, 300]     # 存活率觀察秒數
KEY_SEC    = 60                          # 雜訊%以哪個秒數為準
KEY_SURV   = 70.0                        # 雜訊%：存活率門檻（%）
GAP_SURV   = 80.0                        # 兩單間距：更嚴格的門檻
AMBUSH_SET = [0.5, 0.6, 0.7, 0.8, 0.9, 1.0]
AMBUSH_LOOK = 600
MIN_GAP     = 120
MAX_ENTRY   = 30000
SLIP_TICKS  = 4
REQ_PER_SEC = 6.0
FONT_NAME   = "蘋方-繁"
FONT_SIZE   = 12

def log(*a):
    print(time.strftime("[%H:%M:%S]"), *a, flush=True)

def load_env(path):
    d = {}
    try:
        for line in open(path, encoding="utf-8"):
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1); d[k.strip()] = v.strip().strip('"').strip("'")
    except Exception:
        pass
    return d

_last = [0.0]
def http_get(path, tries=4):
    for i in range(tries):
        gap = 1.0 / REQ_PER_SEC - (time.time() - _last[0])
        if gap > 0:
            time.sleep(gap)
        _last[0] = time.time()
        try:
            req = urllib.request.Request(OKX + path, headers={"User-Agent": "noise-scan/4.0"})
            with urllib.request.urlopen(req, timeout=20) as r:
                j = json.loads(r.read().decode("utf-8"))
            if str(j.get("code")) == "0":
                return j.get("data") or []
            if str(j.get("code")) == "50011":
                time.sleep(1.5 * (i + 1)); continue
            return None
        except Exception:
            time.sleep(1.0 * (i + 1))
    return None

def candidates(sym):
    s = sym.upper()
    for suf in ("USDT", "USDC", "USD"):
        if s.endswith(suf):
            base = s[: -len(suf)]; break
    else:
        base = s
    out = [f"{base}-USDT-SWAP", f"{base}-USD-SWAP"]
    if base == "XAU":
        out.insert(0, "XAUT-USDT-SWAP")
    return out

def fetch_instruments():
    d = http_get("/api/v5/public/instruments?instType=SWAP")
    return {x["instId"]: x for x in d} if d else {}

def resolve(sym, insts):
    for c in candidates(sym):
        if c in insts:
            return c
    return None

def probe_bar(iid):
    for bar in ("1s", "1m"):
        for limit in (300, 100):
            d = http_get(f"/api/v5/market/candles?instId={iid}&bar={bar}&limit={limit}")
            if d and len(d) > 1:
                return bar, (300 if len(d) > 100 else 100)
    return None, 0

def tick_str(t):
    s = f"{t:.10f}".rstrip("0")
    return s + "0" if s.endswith(".") else s

# ==================== 下載 ====================
def fetch_series(iid, bar, page, seconds_back):
    need = seconds_back * 1000
    newest = oldest = None
    chunks = []; total = 0; pages = 0
    while True:
        q = f"/api/v5/market/history-candles?instId={iid}&bar={bar}&limit={page}"
        if oldest:
            q += f"&after={oldest}"
        d = http_get(q)
        if not d:
            break
        chunks.append(np.array([[float(r[0]), float(r[2]), float(r[3]), float(r[4])]
                                for r in d], dtype=np.float64))
        total += len(d); pages += 1
        if newest is None:
            newest = float(d[0][0])
        oldest = d[-1][0]
        if pages % 150 == 0:
            od = datetime.datetime.fromtimestamp(float(oldest) / 1000).strftime("%m-%d %H:%M")
            log(f"    …{total} 根，最舊 {od}")
        if newest - float(oldest) >= need or total > 2_500_000:
            break
    if not chunks:
        return None
    arr = np.concatenate(chunks); del chunks
    arr = arr[np.argsort(arr[:, 0])]
    _, keep = np.unique(arr[:, 0], return_index=True)
    return arr[np.sort(keep)]

def get_data(sym, iid, bar, page, days, use_cache):
    os.makedirs(CACHE_DIR, exist_ok=True)
    cf = os.path.join(CACHE_DIR, f"v4_{sym}_{bar}_{days}d.npy")
    if use_cache and os.path.exists(cf):
        try:
            a = np.load(cf); log(f"  {sym} 讀快取 {len(a)} 根"); return a
        except Exception:
            pass
    log(f"  {sym} 下載 {bar} 約 {days} 天 …")
    a = fetch_series(iid, bar, page, days * 86400)
    if a is None or len(a) < 5000:
        log(f"  {sym} 資料不足，略過"); return None
    np.save(cf, a)
    log(f"  {sym} 完成 {len(a)} 根")
    return a

# ==================== 核心：雜訊存活分析 ====================
def survival(h, l, c, maxsec):
    """量「價格第一次逆向走到 X% 需要多久」。
    多空各算一次再合併 —— 做多怕跌、做空怕漲，兩邊都是雜訊。
    回 firstk[樣本, 距離]，未觸及者為 maxsec+1。"""
    n = len(c) - maxsec - 2
    if n <= 1000:
        return None, 0
    step = max(1, n // MAX_ENTRY)
    idx = np.arange(0, n, step, dtype=np.int64)
    X = np.array(DIST) / 100.0
    outs = []
    for d in ("L", "S"):
        entry = c[idx]
        cur = entry.copy()
        firstk = np.full((len(idx), len(X)), maxsec + 1, dtype=np.int32)
        done = np.zeros((len(idx), len(X)), dtype=bool)
        for k in range(1, maxsec + 1):
            j = idx + k
            if d == "L":
                cur = np.minimum(cur, l[j])
                adv = (entry - cur) / entry          # 做多：跌多少
            else:
                cur = np.maximum(cur, h[j])
                adv = (cur - entry) / entry          # 做空：漲多少
            for q in range(len(X)):
                if done[:, q].all():
                    continue
                hit = (~done[:, q]) & (adv >= X[q])
                if hit.any():
                    firstk[hit, q] = k
                    done[hit, q] = True
            if done.all():
                break
        outs.append(firstk)
    return np.vstack(outs), len(idx) * 2

def surv_table(firstk):
    """回 list of dict：每個距離的各秒數存活率與中位撐多久。"""
    rows = []
    for q, X in enumerate(DIST):
        col = firstk[:, q]
        r = {"dist": X}
        for s in SURV_SEC:
            r[f"s{s}"] = float((col > s).mean() * 100.0)
        alive = col[col <= max(SURV_SEC)]
        r["med"] = float(np.median(alive)) if len(alive) else float("nan")
        r["never"] = float((col > max(SURV_SEC)).mean() * 100.0)
        rows.append(r)
    return rows

def need_dist(rows, sec, want):
    """要達到 want% 存活率，需要離多遠。線性內插；超出網格回最大值。"""
    key = f"s{sec}"
    prev = None
    for r in rows:
        if r[key] >= want:
            if prev is None:
                return r["dist"]
            x0, y0 = prev["dist"], prev[key]
            x1, y1 = r["dist"], r[key]
            if y1 == y0:
                return x1
            return x0 + (want - y0) * (x1 - x0) / (y1 - y0)
        prev = r
    return float(DIST[-1])

def adverse_depth(h, l, c, W=120):
    n = len(c) - W - 1
    if n <= 100:
        return np.array([])
    fmax = sliding_window_view(h[1:], W).max(axis=1)[:n]
    fmin = sliding_window_view(l[1:], W).min(axis=1)[:n]
    c0 = c[:n]
    up = (fmax - c0) / c0 * 100.0
    dn = (c0 - fmin) / c0 * 100.0
    return np.concatenate([up[dn > up], dn[up > dn]])

def _thin(mask, gap):
    idx = np.flatnonzero(mask)
    if len(idx) == 0:
        return 0
    cnt = 1; last = idx[0]
    for i in idx[1:]:
        if i - last >= gap:
            cnt += 1; last = i
    return cnt

def ambush_count(h, l, off_pct, look, gap):
    off = off_pct / 100.0
    if len(l) <= look + 100:
        return 0
    rmin = np.full(len(l), np.inf); rmin[look - 1:] = sliding_window_view(l, look).min(axis=1)
    rmax = np.full(len(h), -np.inf); rmax[look - 1:] = sliding_window_view(h, look).max(axis=1)
    return _thin(h >= rmin * (1 + off), gap) + _thin(l <= rmax * (1 - off), gap)

def pctl(a, q):
    return float(np.percentile(a, q)) if len(a) else float("nan")

def analyse(sym, tick, arr):
    ts, h, l, c = arr[:, 0], arr[:, 1], arr[:, 2], arr[:, 3]
    px = float(np.median(c))
    firstk, nsamp = survival(h, l, c, max(SURV_SEC))
    if firstk is None:
        return None
    rows = surv_table(firstk)
    noise = need_dist(rows, KEY_SEC, KEY_SURV)
    gapd = need_dist(rows, KEY_SEC, GAP_SURV)
    adv = adverse_depth(h, l, c)
    tickp = tick / px * 100.0
    floor3 = tickp * 3.0
    hug = max(noise, floor3)
    slip = tickp * SLIP_TICKS
    return {"sym": sym, "px": px, "tick": tick, "tickp": tickp, "floor3": floor3,
            "noise": noise, "gap": max(gapd, hug + floor3),
            "hug": hug, "slmin": max(gapd, hug + floor3) + hug + slip,
            "p50": pctl(adv, 50), "p60": pctl(adv, 60),
            "p70": pctl(adv, 70), "p80": pctl(adv, 80),
            "amb": {a: ambush_count(h, l, a, AMBUSH_LOOK, MIN_GAP) for a in AMBUSH_SET},
            "rows": rows, "nsamp": nsamp,
            "days": (ts[-1] - ts[0]) / 86400000.0, "bars": len(c)}

# ==================== Excel ====================
HEAD = PatternFill("solid", fgColor="1F3864")

def style(ws, ncol, widths):
    for i in range(1, ncol + 1):
        cl = ws.cell(row=1, column=i)
        cl.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color="FFFFFF")
        cl.fill = HEAD
        cl.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cl.column_letter].width = widths[i - 1]
    for row in ws.iter_rows(min_row=2):
        for cl in row:
            cl.font = Font(name=FONT_NAME, size=FONT_SIZE)
            cl.alignment = Alignment(horizontal="center")

def write_xlsx(res, path, days, bar):
    res = sorted(res, key=lambda r: r["noise"])          # 雜訊由小到大
    wb = Workbook(); ws = wb.active; ws.title = "雜訊總表"
    cols = ["幣種", "最新價", "tick", "一檔%", "三檔地板%", "雜訊%",
            "逆行P50%", "逆行P60%", "逆行P70%", "逆行P80%",
            "埋伏0.5%\n成交次數", "埋伏0.6%\n成交次數", "埋伏0.7%\n成交次數",
            "埋伏0.8%\n成交次數", "埋伏0.9%\n成交次數", "埋伏1.0%\n成交次數",
            "建議兩單間距%", "建議緊貼度%", "SL下限%"]
    w = [11, 12, 11, 9, 10, 9, 10, 10, 10, 10, 11, 11, 11, 11, 11, 11, 13, 13, 11]
    ws.append(cols)
    for r in res:
        ws.append([r["sym"], round(r["px"], 8), tick_str(r["tick"]),
                   round(r["tickp"], 4), round(r["floor3"], 4), round(r["noise"], 3),
                   round(r["p50"], 3), round(r["p60"], 3), round(r["p70"], 3), round(r["p80"], 3)]
                  + [r["amb"][a] for a in AMBUSH_SET]
                  + [round(r["gap"], 3), round(r["hug"], 3), round(r["slmin"], 3)])
    style(ws, len(cols), w)
    ws.freeze_panes = "B2"

    ws2 = wb.create_sheet("雜訊存活表")
    h2 = ["幣種", "距離%"] + [f"撐過{s}秒\n存活率%" for s in SURV_SEC] + ["被掃到的中位秒數"]
    ws2.append(h2)
    for r in res:
        for x in r["rows"]:
            ws2.append([r["sym"], x["dist"]] + [round(x[f"s{s}"], 1) for s in SURV_SEC]
                       + [round(x["med"], 1) if x["med"] == x["med"] else ""])
    style(ws2, len(h2), [11, 9] + [11] * len(SURV_SEC) + [16])
    ws2.freeze_panes = "C2"

    ws3 = wb.create_sheet("怎麼看")
    for ln in [
        f"資料：OKX {bar} K線 約 {days} 天",
        "",
        "■ 雜訊% ＝ 讓部位能撐過 60 秒、存活率達 70% 所需的距離。",
        "　白話：停損要離現價多遠，才不會被日常來回震盪掃掉。",
        "　這是整張表的核心，其他欄位都是它的佐證或延伸。",
        "",
        "■ 怎麼驗證這個數字可不可信：翻到「雜訊存活表」，找你的幣、找 0.1% 那一列，",
        "　看「撐過10秒存活率」。你實盤緊貼 0.1% 兩秒被掃，如果那個數字很低，",
        "　代表模型和你的實盤吻合，這張表可以信。如果很高，代表模型還是錯的，告訴我。",
        "",
        "■ 建議緊貼度% ＝ 雜訊%（若低於三檔地板則取地板，因為引擎設不出更小的）。",
        "■ 建議兩單間距% ＝ 撐過 60 秒存活率達 80% 的距離，比緊貼更嚴格，",
        "　因為 B 單一旦被雜訊觸發就多付一次手續費，門檻要更高。",
        "",
        "■ SL下限% ＝ 建議間距 + 建議緊貼 + 滑價緩衝。",
        "　★ SL 必須大於這個數字。否則 B 觸發時 A 的緊貼落點會等於或超過靜態SL，",
        "　　A 一步都動不了，直接吃滿 SL 加滑價。2026-09-22 WIF 那場就是這樣虧 1.244%。",
        "　這是下限不是建議值，實際設定請再留餘裕。",
        "",
        "■ 逆行P50~P80 ＝ 在「這一單最後會賺」的前提下，中途往虧損方向最多走多少。",
        "　用來交叉檢查間距：間距若小於 P50，代表超過一半的好單會被 B 誤觸發。",
        "",
        "■ 埋伏X%成交次數 ＝ 這段期間內，埋伏在 X% 外的單子會被吃到幾次。",
        "　直接告訴你每個幣、每個埋伏率，可以打幾場戰役。次數太少代表要等很久。",
    ]:
        ws3.append([ln])
    for row in ws3.iter_rows():
        for cl in row:
            cl.font = Font(name=FONT_NAME, size=FONT_SIZE)
    ws3.column_dimensions["A"].width = 100
    wb.save(path)

def send_mail(path, name, res, days, bar):
    env = load_env(os.path.join(BASE, ".env"))
    user = env.get("GMAIL_USER"); pwd = (env.get("GMAIL_APP_PASSWORD") or "").replace(" ", "")
    to = env.get("REPORT_TO") or env.get("REPORT_EMAIL_TO") or user
    if not user or not pwd:
        log(f"！ .env 未設定 Gmail，檔案在 {path}"); return
    m = EmailMessage()
    m["Subject"] = f"OKX 雜訊存活分析 v4 {bar} {days}天（{len(res)} 個幣）"
    m["From"] = user; m["To"] = to
    b = [f"{'幣種':<10}{'雜訊%':>8}{'間距%':>8}{'緊貼%':>8}{'SL下限%':>9}"]
    for r in sorted(res, key=lambda x: x["noise"]):
        b.append(f"{r['sym']:<10}{r['noise']:>8.3f}{r['gap']:>8.3f}{r['hug']:>8.3f}{r['slmin']:>9.3f}")
    b += ["", "雜訊% = 撐過 60 秒、存活率 70% 所需的距離。",
          "驗證法：附件「雜訊存活表」找 0.1% 那一列的 10 秒存活率，",
          "應該很低才對得上你實盤兩秒被掃的經驗。"]
    m.set_content("\n".join(b))
    m.add_attachment(open(path, "rb").read(), maintype="application",
                     subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     filename=name)
    with smtplib.SMTP_SSL("smtp.gmail.com", 465) as s:
        s.login(user, pwd); s.send_message(m)
    log(f"已寄送至 {to}")

def do_probe():
    insts = fetch_instruments()
    if not insts:
        log("！ 無法取得商品清單"); return 1
    for sym in SYMBOLS:
        iid = resolve(sym, insts)
        if not iid:
            print(f"{sym:<11}✗ 找不到"); continue
        tick = float(insts[iid]["tickSz"])
        d = http_get(f"/api/v5/market/candles?instId={iid}&bar=1m&limit=1")
        px = float(d[0][4]) if d else 0.0
        bar, _ = probe_bar(iid)
        print(f"{sym:<11}{iid:<20}{tick_str(tick):<13}{tick/px*100 if px else 0:<9.4f}{bar or '✗'}")
    return 0

def do_run(days, nomail, only, use_cache):
    t0 = time.time()
    insts = fetch_instruments()
    if not insts:
        log("！ 無法取得商品清單"); return 1
    syms = [s for s in SYMBOLS if (not only or s.upper() in only)]
    res = []; bar_used = "1s"
    for n, sym in enumerate(syms, 1):
        log(f"[{n}/{len(syms)}] {sym}")
        iid = resolve(sym, insts)
        if not iid:
            continue
        bar, page = probe_bar(iid)
        if not bar:
            continue
        bar_used = bar
        arr = get_data(sym, iid, bar, page, days, use_cache)
        if arr is None:
            continue
        try:
            r = analyse(sym, float(insts[iid]["tickSz"]), arr)
            if not r:
                log(f"  {sym} 樣本不足"); continue
            res.append(r)
            s10 = [x for x in r["rows"] if abs(x["dist"] - 0.10) < 1e-9]
            chk = f"｜0.1%撐過10秒 {s10[0]['s10']:.0f}%" if s10 else ""
            log(f"  → 雜訊 {r['noise']:.3f}%｜間距 {r['gap']:.3f}%｜緊貼 {r['hug']:.3f}%"
                f"｜SL下限 {r['slmin']:.3f}%{chk}")
        except Exception as e:
            log(f"  {sym} 分析失敗 {type(e).__name__}: {e}")
    if not res:
        log("！ 沒有任何幣種分析成功"); return 1
    os.makedirs(OUT_DIR, exist_ok=True)
    day = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    name = f"OKX.noise.v4.{bar_used}.{days}d.{day}.xlsx"
    path = os.path.join(OUT_DIR, name)
    write_xlsx(res, path, days, bar_used)
    log(f"已產生 {path}")
    if not nomail:
        try:
            send_mail(path, name, res, days, bar_used)
        except Exception as e:
            log(f"！ 寄信失敗 {type(e).__name__}: {e}｜檔案仍在 {path}")
    log(f"完成，耗時 {(time.time()-t0)/60:.1f} 分鐘")
    return 0

def main():
    a = sys.argv[1:]
    mode = a[0] if a else "probe"
    days = 7; nomail = False; only = set(); use_cache = False
    for i, x in enumerate(a):
        if x == "--days" and i + 1 < len(a):
            days = max(1, min(30, int(a[i + 1])))
        elif x == "--no-mail":
            nomail = True
        elif x == "--cache":
            use_cache = True
        elif x == "--fresh":
            use_cache = False
        elif x == "--only" and i + 1 < len(a):
            only = {s.strip().upper() for s in a[i + 1].split(",") if s.strip()}
        elif x == "--rate" and i + 1 < len(a):
            globals()["REQ_PER_SEC"] = max(1.0, min(9.0, float(a[i + 1])))
    if mode == "probe":
        return do_probe()
    if mode == "run":
        return do_run(days, nomail, only, use_cache)
    print(__doc__); return 1

if __name__ == "__main__":
    sys.exit(main())
